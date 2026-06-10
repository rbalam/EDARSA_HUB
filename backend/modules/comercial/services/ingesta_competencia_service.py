"""
Ingesta de Competencia - Servicio
=================================
Alimenta datos de competencia (competidor/categoria/producto/precio) cuando faltan,
desde un adjunto (Excel/CSV directo; PDF/imagen/Word via IA Gemini) o un link (scraping
+ IA). Archiva el adjunto en object storage, guarda las filas EXTRAIDAS en staging
(dbo.Comercial_Ingesta_Competencia) para VALIDACION HUMANA y, al confirmar, inserta en
las tablas canonicas Comercial_Competidores + Comercial_CompetidoresMenuItems.

SQL-First, NO-LIVE, sin hardcode de datos de negocio.
"""
import os
import io
import re
import csv
import json
import uuid
import logging
import tempfile
from typing import Optional, List, Dict, Any, Tuple

import requests

from core.db import execute_sql_query_params
from core.server_registry import EDARSAHUB_CONFIG
from core.object_storage import put_object, get_object, APP_NAME

logger = logging.getLogger(__name__)

MODELO_IA = "gemini-2.5-flash"

# Mapeo de extension/mime -> categoria de origen
EXT_PARSE_DIRECTO = {"xlsx", "xls", "csv"}
EXT_IA_FILE = {"pdf", "png", "jpg", "jpeg", "webp"}
EXT_IA_TEXT = {"docx", "doc", "txt"}

MIME = {
    "pdf": "application/pdf", "png": "image/png", "jpg": "image/jpeg",
    "jpeg": "image/jpeg", "webp": "image/webp", "csv": "text/csv", "txt": "text/plain",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "doc": "application/msword",
}

PROMPT_EXTRACCION = (
    "Eres un extractor de datos de menus de restaurantes competidores. "
    "Del contenido proporcionado, extrae una lista de productos con su precio. "
    "Devuelve EXCLUSIVAMENTE un arreglo JSON valido (sin texto adicional, sin markdown) "
    "donde cada elemento tenga las claves: "
    '"competidor" (nombre del restaurante competidor si aparece, si no usa cadena vacia), '
    '"categoria" (categoria/seccion del menu del producto), '
    '"producto" (nombre del producto), '
    '"precio" (numero decimal, solo el valor sin simbolos), '
    '"moneda" (codigo ISO, por defecto "MXN"). '
    "Si no hay precio claro para un producto, omitelo. No inventes datos."
)


def _conn() -> Tuple:
    return (
        EDARSAHUB_CONFIG['host'], EDARSAHUB_CONFIG['port'], EDARSAHUB_CONFIG['database'],
        EDARSAHUB_CONFIG['username'], EDARSAHUB_CONFIG['password'],
    )


def _q(query: str, params: tuple = ()) -> List[Dict]:
    return execute_sql_query_params(*_conn(), query, params)


# ---------------------------------------------------------------------------
# Parseo de JSON robusto desde respuesta del LLM
# ---------------------------------------------------------------------------

def _parse_json_array(text: str) -> List[Dict]:
    if not text:
        return []
    t = text.strip()
    # quitar fences ```json ... ```
    t = re.sub(r"^```(?:json)?", "", t).strip()
    t = re.sub(r"```$", "", t).strip()
    # extraer el primer arreglo [...] si viene envuelto
    start = t.find("[")
    end = t.rfind("]")
    if start != -1 and end != -1 and end > start:
        t = t[start:end + 1]
    try:
        data = json.loads(t)
        return data if isinstance(data, list) else []
    except Exception as e:
        logger.warning(f"[INGESTA] No se pudo parsear JSON IA: {e}")
        return []


def _normaliza_filas(raw: List[Dict]) -> List[Dict]:
    """Normaliza y valida las filas extraidas."""
    filas = []
    for r in raw:
        if not isinstance(r, dict):
            continue
        producto = str(r.get("producto") or r.get("nombre") or "").strip()
        precio_raw = r.get("precio")
        try:
            precio = float(str(precio_raw).replace("$", "").replace(",", "").strip())
        except Exception:
            precio = None
        if not producto or precio is None or precio <= 0:
            continue
        filas.append({
            "competidor": str(r.get("competidor") or "").strip(),
            "categoria": str(r.get("categoria") or r.get("seccion") or "").strip(),
            "producto": producto,
            "precio": round(precio, 2),
            "moneda": (str(r.get("moneda") or "MXN").strip().upper() or "MXN")[:3],
        })
    return filas


# ---------------------------------------------------------------------------
# Extractores
# ---------------------------------------------------------------------------

def parse_tabular(data: bytes, ext: str) -> List[Dict]:
    """Excel/CSV con columnas: competidor, categoria, producto, precio, moneda (flexible)."""
    rows: List[Dict] = []
    if ext == "csv":
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        raw_rows = list(reader)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            headers = [str(h).strip().lower() if h is not None else "" for h in next(it)]
        except StopIteration:
            return []
        raw_rows = [dict(zip(headers, r)) for r in it]

    def pick(d, *names):
        for n in names:
            for k, v in d.items():
                if k and str(k).strip().lower() == n:
                    return v
        return None

    for d in raw_rows:
        rows.append({
            "competidor": pick(d, "competidor", "restaurante", "competitor"),
            "categoria": pick(d, "categoria", "categoría", "seccion", "sección", "category"),
            "producto": pick(d, "producto", "platillo", "nombre", "product", "item"),
            "precio": pick(d, "precio", "price", "importe"),
            "moneda": pick(d, "moneda", "currency"),
        })
    return _normaliza_filas(rows)


async def _ia_extraer(user_text: str, file_path: Optional[str] = None, mime: Optional[str] = None) -> List[Dict]:
    from emergentintegrations.llm.chat import LlmChat, UserMessage, FileContentWithMimeType
    key = os.environ.get("EMERGENT_LLM_KEY", "")
    if not key:
        raise RuntimeError("EMERGENT_LLM_KEY no configurada")
    chat = LlmChat(
        api_key=key,
        session_id=f"ingesta-{uuid.uuid4()}",
        system_message=PROMPT_EXTRACCION,
    ).with_model("gemini", MODELO_IA)
    file_contents = None
    if file_path and mime:
        file_contents = [FileContentWithMimeType(file_path=file_path, mime_type=mime)]
    msg = UserMessage(text=user_text, file_contents=file_contents) if file_contents \
        else UserMessage(text=user_text)
    resp = await chat.send_message(msg)
    return _normaliza_filas(_parse_json_array(resp))


def _docx_a_texto(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    partes = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for tbl in doc.tables:
        for row in tbl.rows:
            partes.append(" | ".join(c.text for c in row.cells))
    return "\n".join(partes)


def _scrape_link(url: str) -> str:
    from bs4 import BeautifulSoup
    headers = {"User-Agent": "Mozilla/5.0 (compatible; EdarsaBot/1.0)"}
    resp = requests.get(url, headers=headers, timeout=25)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    return "\n".join(lines)[:20000]


# ---------------------------------------------------------------------------
# Crear ingesta (archiva + extrae + guarda staging)
# ---------------------------------------------------------------------------

async def crear_ingesta_archivo(
    empresa_id: int, unidades_csv: str, nombre_archivo: str,
    content_type: str, data: bytes, usuario: str,
) -> Dict:
    ext = (nombre_archivo.rsplit(".", 1)[-1] if "." in nombre_archivo else "bin").lower()
    ingesta_id = str(uuid.uuid4())

    # Archivar SIEMPRE en object storage (auditoria)
    storage_path = f"{APP_NAME}/ingesta-competencia/{empresa_id}/{ingesta_id}.{ext}"
    archivo_path, archivo_tam = None, None
    try:
        result = put_object(storage_path, data, content_type or MIME.get(ext, "application/octet-stream"))
        archivo_path = result.get("path", storage_path)
        archivo_tam = result.get("size", len(data))
    except Exception as e:
        logger.error(f"[INGESTA] Error archivando adjunto: {e}")

    # Extraer segun tipo
    origen, metodo, modelo, mensaje, filas = "OTRO", "PLANTILLA", None, None, []
    try:
        if ext in EXT_PARSE_DIRECTO:
            origen = "EXCEL" if ext in ("xlsx", "xls") else "CSV"
            metodo = "PLANTILLA"
            filas = parse_tabular(data, ext)
        elif ext in EXT_IA_FILE:
            origen = "PDF" if ext == "pdf" else "IMAGEN"
            metodo, modelo = "IA_GEMINI", MODELO_IA
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}")
            tmp.write(data); tmp.flush(); tmp.close()
            try:
                filas = await _ia_extraer("Extrae los productos y precios de este archivo.",
                                          file_path=tmp.name, mime=MIME.get(ext, content_type))
            finally:
                try: os.unlink(tmp.name)
                except Exception: pass
        elif ext in EXT_IA_TEXT:
            origen = "WORD" if ext in ("docx", "doc") else "OTRO"
            metodo, modelo = "IA_GEMINI", MODELO_IA
            texto = _docx_a_texto(data) if ext in ("docx", "doc") else data.decode("utf-8", errors="replace")
            filas = await _ia_extraer(f"Extrae los productos y precios de este contenido:\n\n{texto[:20000]}")
        else:
            mensaje = f"Tipo de archivo no soportado: .{ext}"
    except Exception as e:
        mensaje = f"Error en extraccion: {e}"
        logger.error(f"[INGESTA] {mensaje}")

    if not filas and not mensaje:
        mensaje = "No se detectaron filas con precio. Revisa el archivo o la plantilla."

    _insert_staging(ingesta_id, empresa_id, unidades_csv, origen, metodo, modelo,
                    archivo_path, nombre_archivo, content_type, archivo_tam, None,
                    filas, mensaje, usuario)
    return obtener_ingesta(ingesta_id)


async def crear_ingesta_link(empresa_id: int, unidades_csv: str, url: str, usuario: str) -> Dict:
    ingesta_id = str(uuid.uuid4())
    mensaje, filas = None, []
    try:
        texto = _scrape_link(url)
        if not texto:
            mensaje = "El sitio no devolvio contenido legible (posible bloqueo de scraping)."
        else:
            filas = await _ia_extraer(
                f"Extrae los productos y precios del menu en este contenido web (fuente: {url}):\n\n{texto}")
    except Exception as e:
        mensaje = f"No se pudo leer el link (posible bloqueo o error): {e}"
        logger.error(f"[INGESTA] link {url}: {e}")

    if not filas and not mensaje:
        mensaje = "No se detectaron productos con precio en el link."

    _insert_staging(ingesta_id, empresa_id, unidades_csv, "LINK", "IA_GEMINI", MODELO_IA,
                    None, None, None, None, url, filas, mensaje, usuario)
    return obtener_ingesta(ingesta_id)


def _insert_staging(ingesta_id, empresa_id, unidades_csv, origen, metodo, modelo,
                    archivo_path, archivo_nombre, content_type, archivo_tam, fuente_url,
                    filas, mensaje, usuario):
    estado = "PENDIENTE" if filas else "ERROR"
    _q(
        """
        INSERT INTO Comercial_Ingesta_Competencia
            (IngestaID, EmpresaID, UnidadesNegocioIDs, Origen, MetodoExtraccion, ModeloIA,
             ArchivoPath, ArchivoNombre, ArchivoContentType, ArchivoTamano, FuenteUrl,
             FilasExtraidas, TotalFilas, Estado, Mensaje, UsuarioCreacion, FechaCreacion)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,SYSDATETIME())
        """,
        (ingesta_id, empresa_id, unidades_csv, origen, metodo, modelo,
         archivo_path, archivo_nombre, content_type, archivo_tam, fuente_url,
         json.dumps(filas, ensure_ascii=False), len(filas), estado, mensaje, usuario),
    )


# ---------------------------------------------------------------------------
# Lectura / edicion / resolucion
# ---------------------------------------------------------------------------

def _row_to_dict(r: Dict) -> Dict:
    return {
        "ingesta_id": str(r.get("IngestaID")),
        "empresa_id": r.get("EmpresaID"),
        "unidades_negocio_ids": r.get("UnidadesNegocioIDs"),
        "origen": r.get("Origen"),
        "metodo_extraccion": r.get("MetodoExtraccion"),
        "modelo_ia": r.get("ModeloIA"),
        "archivo_path": r.get("ArchivoPath"),
        "archivo_nombre": r.get("ArchivoNombre"),
        "archivo_content_type": r.get("ArchivoContentType"),
        "archivo_tamano": r.get("ArchivoTamano"),
        "fuente_url": r.get("FuenteUrl"),
        "filas": json.loads(r["FilasExtraidas"]) if r.get("FilasExtraidas") else [],
        "total_filas": r.get("TotalFilas", 0),
        "estado": r.get("Estado"),
        "mensaje": r.get("Mensaje"),
        "competidores_creados": r.get("CompetidoresCreados"),
        "items_creados": r.get("ItemsCreados"),
        "usuario_creacion": r.get("UsuarioCreacion"),
        "fecha_creacion": str(r.get("FechaCreacion")) if r.get("FechaCreacion") else None,
        "usuario_resolucion": r.get("UsuarioResolucion"),
        "fecha_resolucion": str(r.get("FechaResolucion")) if r.get("FechaResolucion") else None,
    }


def listar_ingestas(empresa_id: Optional[int], estado: Optional[str]) -> List[Dict]:
    where, params = ["1=1"], []
    if empresa_id is not None:
        where.append("EmpresaID = %s"); params.append(empresa_id)
    if estado:
        where.append("Estado = %s"); params.append(estado)
    rows = _q(
        f"""SELECT IngestaID, EmpresaID, UnidadesNegocioIDs, Origen, MetodoExtraccion, ModeloIA,
                   ArchivoPath, ArchivoNombre, ArchivoContentType, ArchivoTamano, FuenteUrl,
                   NULL AS FilasExtraidas, TotalFilas, Estado, Mensaje, CompetidoresCreados,
                   ItemsCreados, UsuarioCreacion, FechaCreacion, UsuarioResolucion, FechaResolucion
            FROM Comercial_Ingesta_Competencia
            WHERE {' AND '.join(where)} ORDER BY FechaCreacion DESC""",
        tuple(params),
    )
    return [_row_to_dict(r) for r in rows]


def obtener_ingesta(ingesta_id: str) -> Optional[Dict]:
    rows = _q("SELECT * FROM Comercial_Ingesta_Competencia WHERE IngestaID = %s", (ingesta_id,))
    return _row_to_dict(rows[0]) if rows else None


def actualizar_filas(ingesta_id: str, filas: List[Dict]) -> Optional[Dict]:
    ing = obtener_ingesta(ingesta_id)
    if not ing or ing["estado"] not in ("PENDIENTE", "ERROR"):
        return ing
    norm = _normaliza_filas(filas)
    _q(
        """UPDATE Comercial_Ingesta_Competencia
           SET FilasExtraidas = %s, TotalFilas = %s,
               Estado = CASE WHEN %s > 0 THEN 'PENDIENTE' ELSE Estado END
           WHERE IngestaID = %s""",
        (json.dumps(norm, ensure_ascii=False), len(norm), len(norm), ingesta_id),
    )
    return obtener_ingesta(ingesta_id)


def rechazar_ingesta(ingesta_id: str, usuario: str) -> Optional[Dict]:
    _q(
        """UPDATE Comercial_Ingesta_Competencia
           SET Estado = 'RECHAZADO', UsuarioResolucion = %s, FechaResolucion = SYSDATETIME()
           WHERE IngestaID = %s AND Estado IN ('PENDIENTE','ERROR')""",
        (usuario, ingesta_id),
    )
    return obtener_ingesta(ingesta_id)


def _find_or_create_competidor(empresa_id: int, unidad_id: int, nombre: str, usuario: str) -> str:
    rows = _q(
        """SELECT TOP 1 CompetidorID FROM Comercial_Competidores
           WHERE EmpresaID = %s AND NombreCompetidor = %s AND Activo = 1""",
        (empresa_id, nombre),
    )
    if rows:
        return str(rows[0]["CompetidorID"])
    nuevo = str(uuid.uuid4())
    _q(
        """INSERT INTO Comercial_Competidores
            (CompetidorID, EmpresaID, UnidadNegocioID, NombreCompetidor, Pais,
             EsCompetenciaDirecta, EsBenchmarkAspiracional, Prioridad, Activo,
             FechaCreacion, UsuarioCreacion)
           VALUES (%s,%s,%s,%s,'Mexico',0,0,3,1,SYSDATETIME(),%s)""",
        (nuevo, empresa_id, unidad_id, nombre, usuario),
    )
    return nuevo


def confirmar_ingesta(ingesta_id: str, usuario: str) -> Optional[Dict]:
    ing = obtener_ingesta(ingesta_id)
    if not ing:
        return None
    if ing["estado"] == "CONFIRMADO":
        return ing
    filas = ing["filas"]
    if not filas:
        return ing

    # Unidades destino (CSV de empresa_ids); si vacio, usar la empresa de la ingesta
    unidades = [u.strip() for u in (ing.get("unidades_negocio_ids") or "").split(",") if u.strip()]
    empresas_destino = [int(u) for u in unidades] if unidades else [int(ing["empresa_id"])]

    metodo_obtencion = "IMPORTACION_EXCEL" if ing["metodo_extraccion"] == "PLANTILLA" else "IA_WEB_PUBLICO"
    confianza = "ALTA" if ing["metodo_extraccion"] == "PLANTILLA" else "MEDIA"
    es_ia = 0 if ing["metodo_extraccion"] == "PLANTILLA" else 1
    fuente_url = ing.get("fuente_url")
    nombre_default = ing.get("archivo_nombre") or (fuente_url or "Ingesta")

    competidores_creados, items_creados = set(), 0
    for emp in empresas_destino:
        # competidor por fila (usa el nombre de la fila o un default por archivo)
        for f in filas:
            nombre_comp = (f.get("competidor") or "").strip() or f"Competencia ({nombre_default})"[:200]
            comp_id = _find_or_create_competidor(emp, emp, nombre_comp, usuario)
            competidores_creados.add(comp_id)
            _q(
                """INSERT INTO Comercial_CompetidoresMenuItems
                    (CompetidorMenuItemID, CompetidorID, NombreProductoCompetidor,
                     CategoriaCompetidor, Precio, Moneda, FuenteUrl, FechaConsulta,
                     MetodoObtencion, ConfianzaDato, EsDatoManual, EsDatoIA, Activo,
                     FechaCreacion, UsuarioCreacion)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,CAST(SYSDATETIME() AS DATE),%s,%s,%s,%s,1,SYSDATETIME(),%s)""",
                (str(uuid.uuid4()), comp_id, f["producto"][:255], (f.get("categoria") or None),
                 f["precio"], f.get("moneda", "MXN"), fuente_url, metodo_obtencion, confianza,
                 0 if es_ia else 1, es_ia, usuario),
            )
            items_creados += 1

    _q(
        """UPDATE Comercial_Ingesta_Competencia
           SET Estado = 'CONFIRMADO', CompetidoresCreados = %s, ItemsCreados = %s,
               UsuarioResolucion = %s, FechaResolucion = SYSDATETIME()
           WHERE IngestaID = %s""",
        (len(competidores_creados), items_creados, usuario, ingesta_id),
    )
    logger.info(f"[INGESTA] {ingesta_id} confirmada: {len(competidores_creados)} competidores, {items_creados} items")
    return obtener_ingesta(ingesta_id)


def generar_plantilla_excel() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Competencia"
    ws.append(["competidor", "categoria", "producto", "precio", "moneda"])
    ws.append(["Restaurante Ejemplo", "Carnes", "Rib Eye 400g", 480.00, "MXN"])
    ws.append(["Restaurante Ejemplo", "Entradas", "Guacamole", 150.00, "MXN"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
