#!/usr/bin/env python3
"""Genera documento Word con diagnóstico SQL Server"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Configurar estilos
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# ========== TÍTULO ==========
title = doc.add_heading('DIAGNÓSTICO: SQL Server Local para APIs de Ventas del Día', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph('ORIGEN y 130° QUERÉTARO')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True
subtitle.runs[0].font.size = Pt(14)

doc.add_paragraph()
info = doc.add_paragraph()
info.add_run('Fecha: ').bold = True
info.add_run('Diciembre 2025')
info.add_run('\nEstado: ').bold = True
run = info.add_run('REQUIERE INTERVENCIÓN EN SITIO')
run.font.color.rgb = RGBColor(192, 0, 0)
run.bold = True

# ========== ESTADO ACTUAL ==========
doc.add_heading('ESTADO ACTUAL VERIFICADO', level=1)

table1 = doc.add_table(rows=5, cols=3)
table1.style = 'Table Grid'
table1.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['Componente', 'Estado', 'Evidencia']
row = table1.rows[0]
for i, header in enumerate(headers):
    row.cells[i].text = header
    row.cells[i].paragraphs[0].runs[0].bold = True

data = [
    ['API ORIGEN (puerto 8000)', '✅ Activa', '{"status":"API EDARSA funcionando"}'],
    ['API 130QRO (puerto 8001)', '✅ Activa', '{"status":"API EDARSA funcionando"}'],
    ['Query SQL ORIGEN', '❌ Timeout', 'Query tarda >10s y no responde'],
    ['Query SQL 130QRO', '❌ Timeout', 'Query tarda >10s y no responde'],
]

for i, row_data in enumerate(data):
    row = table1.rows[i + 1]
    for j, cell_data in enumerate(row_data):
        row.cells[j].text = cell_data

doc.add_paragraph()
conclusion = doc.add_paragraph()
conclusion.add_run('Conclusión: ').bold = True
conclusion.add_run('Las APIs intermedias funcionan, pero NO pueden conectarse a los SQL Server de las sucursales.')

# ========== ARQUITECTURA ==========
doc.add_heading('ARQUITECTURA ACTUAL', level=1)

arch = doc.add_paragraph()
arch.add_run('''
┌─────────────────┐     HTTP      ┌─────────────────────────┐     SQL      ┌──────────────────┐
│  EDARSA HUB     │──────────────▶│  API Intermedia (Nube)  │──────────────▶│  SQL Server      │
│  (Este servidor)│               │  54.39.104.176          │              │  (Sucursal local)│
└─────────────────┘               └─────────────────────────┘              └──────────────────┘
                                         ✅ OK                                  ❌ FALLA
''').font.name = 'Consolas'

doc.add_paragraph()
doc.add_paragraph('APIs Intermedias Configuradas:').runs[0].bold = True

table2 = doc.add_table(rows=3, cols=5)
table2.style = 'Table Grid'

headers2 = ['Sucursal', 'URL API', 'Puerto', 'Estado HTTP', 'Estado SQL']
for i, h in enumerate(headers2):
    table2.rows[0].cells[i].text = h
    table2.rows[0].cells[i].paragraphs[0].runs[0].bold = True

table2.rows[1].cells[0].text = 'ORIGEN'
table2.rows[1].cells[1].text = 'http://54.39.104.176:8000/query'
table2.rows[1].cells[2].text = '8000'
table2.rows[1].cells[3].text = '✅ Responde'
table2.rows[1].cells[4].text = '❌ SQL timeout'

table2.rows[2].cells[0].text = '130QRO'
table2.rows[2].cells[1].text = 'http://54.39.104.176:8001/query'
table2.rows[2].cells[2].text = '8001'
table2.rows[2].cells[3].text = '✅ Responde'
table2.rows[2].cells[4].text = '❌ SQL timeout'

# ========== CAUSA RAÍZ ==========
doc.add_heading('CAUSA RAÍZ PROBABLE', level=1)

p = doc.add_paragraph()
p.add_run('El problema NO está en EDARSA HUB ni en las APIs intermedias (ambas funcionan).\n')
p.add_run('El problema está en ').bold = False
p.add_run('la conectividad entre las APIs intermedias y el SQL Server en cada sucursal').bold = True
p.add_run('.')

doc.add_paragraph('Posibles causas:').runs[0].bold = True
causas = [
    'SQL Server detenido en la sucursal',
    'Firewall bloqueando puerto SQL (1433 o el configurado)',
    'Red/VPN entre servidor en nube y sucursales no disponible',
    'Credenciales de SQL incorrectas en la configuración de las APIs',
    'TCP/IP deshabilitado en SQL Server Configuration Manager',
    'Puertos dinámicos en lugar de puerto fijo',
]
for causa in causas:
    doc.add_paragraph(causa, style='List Number')

# ========== DIAGNÓSTICO ==========
doc.add_heading('DIAGNÓSTICO A EJECUTAR EN CADA SUCURSAL', level=1)

# Paso 1
doc.add_heading('1. VERIFICAR SERVICIO SQL SERVER', level=2)
code1 = doc.add_paragraph()
code1.add_run('# En la máquina de la sucursal (Windows)\n').font.name = 'Consolas'
code1.add_run('Get-Service -Name "MSSQL*" | Format-Table Name, Status, StartType\n\n').font.name = 'Consolas'
code1.add_run('# Si está detenido:\n').font.name = 'Consolas'
code1.add_run('Start-Service -Name "MSSQLSERVER"').font.name = 'Consolas'

# Paso 2
doc.add_heading('2. VERIFICAR BASE DE DATOS', level=2)
code2 = doc.add_paragraph()
code2.add_run('-- Conectarse localmente con SSMS y ejecutar:\n').font.name = 'Consolas'
code2.add_run("SELECT name, state_desc FROM sys.databases WHERE name = 'ManagementPro'\n").font.name = 'Consolas'
code2.add_run('-- Debe mostrar: ONLINE').font.name = 'Consolas'

# Paso 3
doc.add_heading('3. VERIFICAR CONFIGURACIÓN DE RED SQL', level=2)
pasos_red = [
    'Abrir "SQL Server Configuration Manager"',
    'Ir a "SQL Server Network Configuration" > "Protocols for [INSTANCIA]"',
    'Verificar que TCP/IP esté "Enabled"',
    'Doble clic en TCP/IP > pestaña "IP Addresses"',
    'En "IPAll": TCP Dynamic Ports = VACÍO, TCP Port = 1433',
    'Reiniciar servicio SQL Server',
]
for paso in pasos_red:
    doc.add_paragraph(paso, style='List Number')

# Paso 4
doc.add_heading('4. VERIFICAR FIREWALL (Windows)', level=2)
code4 = doc.add_paragraph()
code4.add_run('# Verificar regla existente\n').font.name = 'Consolas'
code4.add_run('Get-NetFirewallRule -DisplayName "*SQL*" | Format-Table Name, Enabled, Direction\n\n').font.name = 'Consolas'
code4.add_run('# Crear regla si no existe\n').font.name = 'Consolas'
code4.add_run('New-NetFirewallRule -DisplayName "SQL Server" -Direction Inbound -Protocol TCP -LocalPort 1433 -Action Allow').font.name = 'Consolas'

# Paso 5
doc.add_heading('5. PRUEBA DE CONECTIVIDAD DESDE SERVIDOR API', level=2)
doc.add_paragraph('Desde el servidor donde corren las APIs (54.39.104.176):')
code5 = doc.add_paragraph()
code5.add_run('# Para ORIGEN\n').font.name = 'Consolas'
code5.add_run('Test-NetConnection -ComputerName [IP_SUCURSAL_ORIGEN] -Port 1433\n\n').font.name = 'Consolas'
code5.add_run('# Para 130QRO\n').font.name = 'Consolas'
code5.add_run('Test-NetConnection -ComputerName [IP_SUCURSAL_QRO] -Port 1433\n\n').font.name = 'Consolas'
code5.add_run('# Debe devolver: TcpTestSucceeded = True').font.name = 'Consolas'

# Paso 6
doc.add_heading('6. VERIFICAR CREDENCIALES SQL', level=2)
code6 = doc.add_paragraph()
code6.add_run("SELECT name, is_disabled FROM sys.sql_logins WHERE name = 'usuario_api'\n").font.name = 'Consolas'
code6.add_run('-- is_disabled debe ser 0\n\n').font.name = 'Consolas'
code6.add_run("EXEC xp_loginconfig 'login mode'\n").font.name = 'Consolas'
code6.add_run('-- Debe mostrar: Mixed').font.name = 'Consolas'

# Paso 7
doc.add_heading('7. VERIFICAR SQL SERVER BROWSER (si es instancia nombrada)', level=2)
code7 = doc.add_paragraph()
code7.add_run('Get-Service -Name "SQLBrowser" | Format-Table Name, Status\n\n').font.name = 'Consolas'
code7.add_run('# Si está detenido:\n').font.name = 'Consolas'
code7.add_run('Start-Service -Name "SQLBrowser"\n').font.name = 'Consolas'
code7.add_run('Set-Service -Name "SQLBrowser" -StartupType Automatic').font.name = 'Consolas'

# ========== CONNECTION STRING ==========
doc.add_heading('CONNECTION STRING CORRECTO', level=1)
doc.add_paragraph('Las APIs intermedias deben usar un connection string en este formato:')
code_cs = doc.add_paragraph()
code_cs.add_run('# Para instancia predeterminada:\n').font.name = 'Consolas'
code_cs.add_run('Server=IP_SUCURSAL,1433;Database=ManagementPro;User Id=usuario;Password=***;\n\n').font.name = 'Consolas'
code_cs.add_run('# Para instancia nombrada CON puerto fijo:\n').font.name = 'Consolas'
code_cs.add_run('Server=IP_SUCURSAL,PUERTO;Database=ManagementPro;User Id=usuario;Password=***;\n\n').font.name = 'Consolas'
code_cs.add_run('# EVITAR este formato (problemas con algunos drivers):\n').font.name = 'Consolas'
code_cs.add_run('Server=IP_SUCURSAL\\INSTANCIA,PUERTO;Database=...').font.name = 'Consolas'

# ========== ENTREGABLES ==========
doc.add_heading('ENTREGABLES REQUERIDOS POR SUCURSAL', level=1)
doc.add_paragraph('Para cerrar este ticket, entregar para ORIGEN y 130QRO:').runs[0].bold = True

table3 = doc.add_table(rows=10, cols=3)
table3.style = 'Table Grid'

headers3 = ['#', 'Entregable', 'Valor']
for i, h in enumerate(headers3):
    table3.rows[0].cells[i].text = h
    table3.rows[0].cells[i].paragraphs[0].runs[0].bold = True

entregables = [
    ['1', 'IP del SQL Server', 'ej: 192.168.1.100'],
    ['2', 'Puerto configurado', 'ej: 1433'],
    ['3', 'Nombre de instancia', 'ej: DEFAULT o SQLEXPRESS'],
    ['4', 'Resultado Test-NetConnection', 'TcpTestSucceeded = True/False'],
    ['5', 'Estado servicio SQL Server', 'Running/Stopped'],
    ['6', 'Estado SQL Server Browser', 'Running/Stopped/N/A'],
    ['7', 'Causa raíz identificada', 'ej: "Firewall bloqueaba puerto"'],
    ['8', 'Cambios realizados', 'ej: "Habilitado TCP/IP, puerto fijo 1433"'],
    ['9', 'Evidencia de conexión OK', 'Screenshot o resultado de query'],
]

for i, row_data in enumerate(entregables):
    for j, cell_data in enumerate(row_data):
        table3.rows[i + 1].cells[j].text = cell_data

# ========== PRUEBA FINAL ==========
doc.add_heading('PRUEBA FINAL DESDE EDARSA HUB', level=1)
doc.add_paragraph('Una vez corregido en las sucursales, ejecutar:')
code_final = doc.add_paragraph()
code_final.add_run('curl -X GET "http://54.39.104.176:8000/query?sql=SELECT%201%20as%20test" \\\n').font.name = 'Consolas'
code_final.add_run('  -H "x-api-key: ${API_MPRO_KEY}"\n\n').font.name = 'Consolas'
code_final.add_run('curl -X GET "http://54.39.104.176:8001/query?sql=SELECT%201%20as%20test" \\\n').font.name = 'Consolas'
code_final.add_run('  -H "x-api-key: ${API_MPRO_KEY}"').font.name = 'Consolas'

doc.add_paragraph('Respuesta esperada:').runs[0].bold = True
code_resp = doc.add_paragraph()
code_resp.add_run('{"total_registros": 1, "data": [{"test": 1}]}').font.name = 'Consolas'

# ========== NOTA IMPORTANTE ==========
doc.add_heading('NOTA IMPORTANTE', level=1)
nota = doc.add_paragraph()
run_nota = nota.add_run('Este problema NO se puede resolver desde EDARSA HUB.')
run_nota.bold = True
run_nota.font.color.rgb = RGBColor(192, 0, 0)

doc.add_paragraph('Se requiere:')
reqs = [
    'Acceso físico o remoto a las máquinas de las sucursales',
    'Permisos de administrador en Windows',
    'Acceso a SQL Server Management Studio',
    'Posiblemente acceso al router/firewall de la sucursal',
]
for req in reqs:
    doc.add_paragraph(req, style='List Number')

doc.add_paragraph()
final = doc.add_paragraph()
final.add_run('El código de EDARSA HUB está correcto y listo. Solo falta restablecer la conectividad SQL en las sucursales.').italic = True

# Guardar
doc.save('/app/docs/DIAGNOSTICO_SQL_SUCURSALES.docx')
print("Documento Word generado: /app/docs/DIAGNOSTICO_SQL_SUCURSALES.docx")
