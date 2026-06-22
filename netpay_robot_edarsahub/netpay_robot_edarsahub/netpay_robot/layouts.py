from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any
from openpyxl import load_workbook


TRANSACCIONES_SHEET = 'TRX Tarjeta Presente'
DEPOSITOS_RESUMEN_SHEET = 'Resumen'
DEPOSITOS_DETALLE_SHEET = 'Ventas Tarjeta Presente'

TRANSACCIONES_REQUIRED = [
    'Fecha de TRX', 'Hora de TRX', 'Monto de trx', 'Venta Neta', 'Estatus de trx',
    'Código de respuesta', 'Motivo', 'Company name', 'Sucursal', 'Store ID',
    'Propina', 'Banco', 'Marca', 'Tipo de Venta', 'Tipo de Tarjeta',
    'Order ID', 'Código de Autorización', 'Referencia'
]

DEPOSITOS_RESUMEN_REQUIRED = [
    'Fecha de Movimiento', 'Clave Rastreo', 'Cuenta Depósito', 'Descripción', 'Monto Depósito'
]

DEPOSITOS_DETALLE_REQUIRED = [
    'Fecha de depósito', 'Clave Rastreo', 'Cuenta Depósito', 'Nombre Empresa', 'Sucursal',
    'Store ID', 'Producto', 'Monto Depósito', 'Fecha Trx', 'Hora de Trx', 'Monto de Trx',
    'Venta Neta', 'Propina', 'Comisión Base (%)', 'Comisión Base ($)',
    'Comisión Total ($)', 'IVA Comisiones (16%)', 'Comisiones + IVA',
    'Banco', 'Marca', 'Tipo de Venta', 'Tipo de Tarjeta', 'Código de Autorización',
    'Order ID', 'Referencia'
]


@dataclass
class SheetLayout:
    sheet_name: str
    header_row: int
    headers: list[str]
    missing_required: list[str]


@dataclass
class WorkbookLayoutResult:
    ok: bool
    report_kind: str | None
    layouts: list[SheetLayout]
    errors: list[str]
    warnings: list[str]


def _normalize_header(value: Any) -> str:
    if value is None:
        return ''
    return str(value).replace('\n', ' ').strip()


def _find_header_row(ws, required: list[str], search_first_rows: int = 30) -> tuple[int | None, list[str], list[str]]:
    required_norm = {_normalize_header(x).lower(): x for x in required}
    best_row = None
    best_headers: list[str] = []
    best_score = -1
    for row_idx in range(1, min(ws.max_row, search_first_rows) + 1):
        headers = [_normalize_header(ws.cell(row=row_idx, column=col).value) for col in range(1, ws.max_column + 1)]
        score = sum(1 for h in headers if h.lower() in required_norm)
        if score > best_score:
            best_score = score
            best_row = row_idx
            best_headers = headers
    present = {h.lower() for h in best_headers if h}
    missing = [original for key, original in required_norm.items() if key not in present]
    if best_score == 0:
        return None, [], required
    return best_row, best_headers, missing


def detect_layout(path: str | Path) -> WorkbookLayoutResult:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_names = set(wb.sheetnames)
    errors: list[str] = []
    warnings: list[str] = []
    layouts: list[SheetLayout] = []

    if TRANSACCIONES_SHEET in sheet_names:
        ws = wb[TRANSACCIONES_SHEET]
        row, headers, missing = _find_header_row(ws, TRANSACCIONES_REQUIRED)
        if row is None:
            errors.append(f'No se encontraron encabezados válidos en {TRANSACCIONES_SHEET}.')
        layouts.append(SheetLayout(TRANSACCIONES_SHEET, row or -1, headers, missing))
        if missing:
            errors.append(f'Faltan columnas en {TRANSACCIONES_SHEET}: {missing}')
        return WorkbookLayoutResult(not errors, 'DETALLE_TRANSACCIONES', layouts, errors, warnings)

    has_resumen = DEPOSITOS_RESUMEN_SHEET in sheet_names
    has_detalle = DEPOSITOS_DETALLE_SHEET in sheet_names
    if has_resumen and has_detalle:
        ws_res = wb[DEPOSITOS_RESUMEN_SHEET]
        row_res, headers_res, missing_res = _find_header_row(ws_res, DEPOSITOS_RESUMEN_REQUIRED)
        layouts.append(SheetLayout(DEPOSITOS_RESUMEN_SHEET, row_res or -1, headers_res, missing_res))
        if missing_res:
            errors.append(f'Faltan columnas en {DEPOSITOS_RESUMEN_SHEET}: {missing_res}')

        ws_det = wb[DEPOSITOS_DETALLE_SHEET]
        row_det, headers_det, missing_det = _find_header_row(ws_det, DEPOSITOS_DETALLE_REQUIRED)
        layouts.append(SheetLayout(DEPOSITOS_DETALLE_SHEET, row_det or -1, headers_det, missing_det))
        if missing_det:
            errors.append(f'Faltan columnas en {DEPOSITOS_DETALLE_SHEET}: {missing_det}')
        return WorkbookLayoutResult(not errors, 'DETALLE_DEPOSITOS_MOVIMIENTOS', layouts, errors, warnings)

    errors.append(f'No se reconoció layout NetPay. Hojas encontradas: {wb.sheetnames}')
    return WorkbookLayoutResult(False, None, layouts, errors, warnings)
